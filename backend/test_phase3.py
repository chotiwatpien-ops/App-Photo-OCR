# -*- coding: utf-8 -*-
"""Rider Trips Phase 3: every line of the passenger's fare block, from W38 (Ops 2026-09-21).

What must hold: each row adds up across its own columns — paid + every line − tip = ride fare —
so a customer checking a row with a calculator lands on the slip's own figure. The net is the
same number it always was, only with the tip in a column of its own. Phase 2 stops at W37 and
keeps its look; nothing already delivered changes. The money below is real W38 slips.
"""
import io
import sys

import openpyxl

sys.stdout.reconfigure(encoding="utf-8")
sys.path.insert(0, "backend")
import excel_writer as xw                                          # noqa: E402

ok = True


def check(name, cond):
    global ok
    print(("✅" if cond else "✗"), name)
    ok = ok and cond


def trip(date_="2026-09-15", name="สมชาย", **money):
    t = {"driver_name": name, "trip_date": date_, "trip_time": "14:30",
         "service_type": "Standard Bike", "payment_method": "GRAB PAY", "distance_km": 5.0,
         "pickup_text": "เอ็มสเฟียร์", "dropoff_text": "มหานคร", "pickup_district": "วัฒนา",
         "dropoff_district": "บางรัก", "customer_image": "WK38-สมชาย1.jpg"}
    t.update(money)
    return t


# real slips (read by eye 2026-09-21); bonus as stored = bonus + tip (pipeline.apply_extraction)
SLIPS = {
    "ops": trip(base_fare=23, passenger_paid=29, app_fee=-1, intl_fee=1, passenger_total=27),
    "promo_31153": trip(name="ก", base_fare=95, turbo=5, passenger_paid=90, app_fee=-20,
                        discount=50, passenger_total=120),
    "both_32256": trip(name="ข", base_fare=27, passenger_paid=39, app_fee=-1, discount=3,
                       insurance_fee=-10, passenger_total=31),
    "tolls_31974": trip(name="ค", base_fare=220, passenger_paid=458, app_fee=-20, intl_fee=13,
                        passenger_tolls=-140, passenger_total=285),
    "everything_31420": trip(name="ง", base_fare=358, passenger_paid=425, app_fee=-20,
                             insurance_fee=-5, intl_fee=12, discount=104, other_adj=-50,
                             passenger_total=442),
    "tip_31962": trip(name="จ", base_fare=125, bonus=40, tip=40, passenger_paid=206,
                      app_fee=-20, discount=10, passenger_total=156),
    "old_screen": trip(name="ฉ", base_fare=100, passenger_total=120),
    # the slip shows tip 20 in the extra income but 'คุณได้รับ ฿78' = base 74 + turbo 4
    "tip_outside_net_30298": trip(name="ช", base_fare=74, turbo=4, bonus=0, tip=20,
                                  passenger_paid=144, app_fee=-20, intl_fee=4, passenger_total=100),
}
data = xw.build_fare_lines_workbook(list(SLIPS.values()))
ws = openpyxl.load_workbook(io.BytesIO(data))[xw.LOCATION_SHEET]
H = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
by_name = {ws.cell(r, 1).value: r for r in range(2, ws.max_row + 1)}


def v(key, header):
    return ws.cell(by_name[SLIPS[key]["driver_name"]], H[header]).value


def adds_up(key):
    got = (v(key, "Passenger Fare (THB)") + v(key, "International Fee") + v(key, "Application Fee")
           + v(key, "Discount") + v(key, "Travel Insurance Fee") + v(key, "Passenger Tolls")
           + v(key, "Other Fees") - v(key, "Tip (THB)"))
    return abs(got - v(key, "Ride Fare (THB)")) < 0.01


# --- the layout ---------------------------------------------------------------------------
check("หัวคอลัมน์ตรงกับแบบที่ตกลง A–Y แล้วโซนต่อท้าย",
      [ws.cell(1, c).value for c in range(1, ws.max_column + 1)] == xw.FARE_LINES_HEADERS)
check("A–K ชื่อเดิมตำแหน่งเดิม (ลูกค้าอ้างตัวอักษรได้เหมือนเดิม)",
      xw.FARE_HEADERS[:11] == xw.HEADERS[:11])
check("L–O ชื่อเดิมตำแหน่งเดิม", xw.FARE_HEADERS[11:15] == xw.HEADERS[11:15])
check("ชื่อ Ride Fare ตามที่ Joe ขอ", "Ride Fare (THB)" in H and "Total Passenger Fare" not in str(H))
check("ความกว้างครบทุกคอลัมน์", len(xw.FARE_WIDTHS) == len(xw.FARE_HEADERS))
check("ชีตชื่อเดียวกับ Phase 2 (อะไรที่อ่าน Phase 2 อยู่ หาเจอ)", ws.title == xw.LOCATION_SHEET)
check("LAYOUT เปลี่ยนแล้ว — ไฟล์บน Drive ถูกเขียนใหม่รอบหน้าแน่ ๆ", xw.LAYOUT != "2026-09-17")

# --- formulas point at the right columns ------------------------------------------------------
r = by_name["สมชาย"]
check("J = K+M+N+P (ทิปมีช่องของตัวเองแล้ว)", ws.cell(r, H["Net Earnings (THB)"]).value == f"=K{r}+M{r}+N{r}+P{r}")
check("X = W−K (ค่าบริการ Grab จากค่ารอบ)", ws.cell(r, H["Grab Service Fee (THB)"]).value == f"=W{r}-K{r}")
check("K คือ Base Fare, P คือ Tip, W คือ Ride Fare",
      (H["Base Fare (THB)"], H["Tip (THB)"], H["Ride Fare (THB)"]) == (11, 16, 23))

# --- every row adds up, on real slips ---------------------------------------------------------
for key in ("ops", "promo_31153", "both_32256", "tolls_31974", "everything_31420", "tip_31962"):
    check(f"{key}: ยอดชำระ + ทุกบรรทัด − ทิป = ค่ารอบ", adds_up(key))
check("สลิปของ Ops: ค่าบริการ Grab = 27 − 23 = 4",
      v("ops", "Ride Fare (THB)") - v("ops", "Base Fare (THB)") == 4)
check("#31420 ครบทุกบรรทัดในใบเดียว: อื่นๆ −50 อยู่ใน Other Fees", v("everything_31420", "Other Fees") == -50)

# --- signs as the slip prints them ------------------------------------------------------------
check("L ค่าธุรกรรมต่างประเทศติดลบตามสลิป", v("ops", "International Fee") == -1)
check("L ไม่มีค่า = 0 ไม่ใช่ −0", v("promo_31153", "International Fee") == 0)
check("ค่าแอปติดลบตามสลิป", v("ops", "Application Fee") == -1)
check("ส่วนลดเป็นบวก ค่าประกันเป็นลบ",
      v("both_32256", "Discount") == 3 and v("both_32256", "Travel Insurance Fee") == -10)
check("ค่าทางด่วนฝั่งผู้โดยสารติดลบ", v("tolls_31974", "Passenger Tolls") == -140)

# --- the tip leaves Bonus, the net stays the same -------------------------------------------
check("#31962 Bonus เหลือโบนัสจริง (0) ทิปไปอยู่ช่อง Tip (40)",
      v("tip_31962", "Bonus") == 0 and v("tip_31962", "Tip (THB)") == 40)
t = SLIPS["tip_31962"]
check("#31962 net เท่าเดิม: K+M+N+P = 125+0+0+40 = 165 = สูตรเก่า K+M+N (Bonus รวมทิป)",
      v("tip_31962", "Base Fare (THB)") + v("tip_31962", "Bonus") + v("tip_31962", "Turbo Incentive (THB)")
      + v("tip_31962", "Tip (THB)") == t["base_fare"] + t["bonus"] == 165)

check("❗ #30298 ทิปที่ไม่ได้อยู่ในรายได้สุทธิ: Bonus ไม่ติดลบ",
      v("tip_outside_net_30298", "Bonus") == 0)
check("❗ #30298 net ไม่ขยับ: 74 + 0 + 4 + 0 = 78 ตาม 'คุณได้รับ'",
      v("tip_outside_net_30298", "Base Fare (THB)") + v("tip_outside_net_30298", "Bonus")
      + v("tip_outside_net_30298", "Turbo Incentive (THB)") + v("tip_outside_net_30298", "Tip (THB)") == 78)

# --- the old app screen ------------------------------------------------------------------------
check("จอเก่าไม่มียอดชำระ: Passenger Fare ว่าง ไม่เดาเลข", v("old_screen", "Passenger Fare (THB)") is None)
check("จอเก่ายังมีค่ารอบ", v("old_screen", "Ride Fare (THB)") == 120)

# --- which week goes where ---------------------------------------------------------------------
w36, w37, w38, w39 = (trip("2026-08-31", "ก", base_fare=1), trip("2026-09-08", "ข", base_fare=1),
                      trip("2026-09-14", "ค", base_fare=1), trip("2026-09-21", "ง", base_fare=1))
p2 = openpyxl.load_workbook(io.BytesIO(xw.build_location_workbook([w36, w37, w38, w39])))[xw.LOCATION_SHEET]
p3 = openpyxl.load_workbook(io.BytesIO(xw.build_fare_lines_workbook([w36, w37, w38, w39])))[xw.LOCATION_SHEET]
check("❗ Phase 2 หยุดที่ W37 หน้าตาเดิม (W36+W37 = 2 แถว)", p2.max_row == 3
      and [p2.cell(1, c).value for c in range(1, len(xw.LOCATION_HEADERS) + 1)] == xw.LOCATION_HEADERS)
check("Phase 3 เริ่ม W38 (W38+W39 = 2 แถว)", p3.max_row == 3)
check("ไม่มีแถวไหนอยู่สองไฟล์", {p2.cell(r, 1).value for r in (2, 3)}.isdisjoint({p3.cell(r, 1).value for r in (2, 3)}))
check("❗ ไฟล์แรก (ก่อน W36) ไม่รับแถวใหม่",
      openpyxl.load_workbook(io.BytesIO(xw.build_workbook([w38])))[xw.SHEET].max_row == 1)
check("ยังไม่มีแถว W38: ไม่สร้างไฟล์เปล่า", xw.build_fare_lines_workbook([w36, w37]) is None)
names = [n for n, _ in xw.later_workbooks([w36, w37, w38])]
check("อัปโหลดทั้ง Phase 2 และ Phase 3 ตามลำดับ", names == [xw.LOCATION_FILE, xw.FARE_LINES_FILE])
check("สัปดาห์ที่ยังไม่ถึง Phase 3 อัปโหลดแค่ Phase 2", [n for n, _ in xw.later_workbooks([w37])] == [xw.LOCATION_FILE])

print("\nสรุป:", "ผ่านทั้งหมด ✅" if ok else "มีข้อที่ไม่ผ่าน ✗")
sys.exit(0 if ok else 1)
