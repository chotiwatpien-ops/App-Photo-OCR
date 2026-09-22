# -*- coding: utf-8 -*-
"""A delivered week comes out again in the Phase 3 layout, as its own file (W35, 2026-09-22)."""
import io
import os
import shutil
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
WORK = tempfile.mkdtemp(prefix="pocr-wk3-")
os.environ["PHOTO_OCR_DATA"] = WORK
os.environ.pop("DATABASE_URL", None)
sys.path.insert(0, "backend")

from openpyxl import load_workbook                              # noqa: E402
from sqlalchemy import insert                                   # noqa: E402

import db                                                       # noqa: E402
import excel_writer                                             # noqa: E402
import week_phase3 as wp                                        # noqa: E402

db.init_db()
ok = True


def check(name, cond):
    global ok
    print(("✅" if cond else "✗"), name)
    ok = ok and cond


W35 = ("2026-08-24", "2026-08-30")
j = db.create_job("สมชาย Win", "Trips", *W35, category="2 W Saver")
other = db.create_job("สมชาย Win", "Trips", "2026-08-31", "2026-09-06", category="2 W Saver")
with db.engine.begin() as c:
    base = dict(status="done", committed=1, service_type="Saver Bike", base_fare=30.0, net_earnings=30.0,
                passenger_paid=40.0, passenger_total=35.0, app_fee=1.0, discount=4.0)
    c.execute(insert(db.trips).values(job_id=j, file_name="a.jpg", trip_date="2026-08-25", **base))
    c.execute(insert(db.trips).values(job_id=j, file_name="b.jpg", trip_date="2026-08-24", **base))
    c.execute(insert(db.trips).values(job_id=j, file_name="dup.jpg", trip_date="2026-08-24",
                                      **{**base, "status": "duplicate"}))
    c.execute(insert(db.trips).values(job_id=j, file_name="wait.jpg", trip_date="2026-08-24",
                                      **{**base, "committed": 0}))
    c.execute(insert(db.trips).values(job_id=other, file_name="w36.jpg", trip_date="2026-08-31", **base))

rows = wp.week_rows(*W35)
check("เฉพาะแถวของสัปดาห์นั้นที่อนุมัติแล้ว (ไม่เอาซ้ำ ไม่เอาที่รอคน ไม่เอาสัปดาห์อื่น)",
      sorted(r["file_name"] for r in rows) == ["a.jpg", "b.jpg"])
check("ชื่อไฟล์บอกสัปดาห์และ Phase 3", wp.file_name(W35[0]) == "Rider Trips 2026-W35 Phase 3.xlsx")
out = os.path.join(WORK, "w35.xlsx")
check("รันได้ เขียนไฟล์ได้", wp.main(["--from", W35[0], "--to", W35[1], "--xlsx", out]) == 0)
ws = load_workbook(out).active
head = [c.value for c in ws[1]]
check("❗ หัวคอลัมน์เป็นแบบ Phase 3 แม้วันที่จะก่อน W38",
      head[:len(excel_writer.FARE_LINES_HEADERS)] == excel_writer.FARE_LINES_HEADERS
      and "Total Commission (THB)" in head)
check("สองแถว เรียงตามวันที่", ws.max_row == 3 and ws.cell(2, 1).value <= ws.cell(3, 1).value)
check("ส่วนลดอยู่คอลัมน์ของตัวเอง", ws.cell(2, head.index("Discount") + 1).value in (4, 4.0))

shutil.rmtree(WORK, ignore_errors=True)
print("\nสรุป:", "ผ่านทั้งหมด ✅" if ok else "มีข้อที่ไม่ผ่าน ✗")
sys.exit(0 if ok else 1)
