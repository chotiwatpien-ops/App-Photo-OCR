# -*- coding: utf-8 -*-
"""What is wrong with one delivered week, row by row — read only.

Ops asked for W34 and W35 "as clean and as matched as it can be" (Fiat 2026-09-22). The Phase 3
re-issue of W34 matched the delivered file row for row, but the week itself carried faults from
the days it was read: 58 pairs of rows equal in every visible column, most of them under two
different riders (฿9,055 of fares), nine rows of one rider with no base fare, 499 picture names
used twice. The Excel file has no booking code, so none of that could be settled from it. This
reads the database, where the code is, and lists each kind of fault on its own sheet. Nothing is
changed; what to do with each list is decided afterwards.

    python week_audit.py --from 2026-08-17 --to 2026-08-23                # summary
    python week_audit.py --from 2026-08-17 --to 2026-08-23 --to-drive     # + Exports/_ตรวจสัปดาห์/
"""
import argparse
import io
import sys
from collections import Counter, defaultdict

import db
from ingest import week_label
from reread_passenger import gap
from week_phase3 import week_rows

GROUP_OF = {"Saver Bike": "2 W Saver", "Standard Bike": "2 W Standard",
            "Standard Car": "4 W Standard", "Saver Car": "4 W Standard"}
DRIVE_DIR = "_ตรวจสัปดาห์"
COLS = ["id", "ไรเดอร์", "กลุ่ม", "Service Type", "วันที่", "เวลา", "รหัสการจอง", "ค่ารอบ", "โบนัส", "Turbo",
        "ทิป", "ค่าทางด่วนคืน", "คุณได้รับ", "ผู้โดยสารจ่าย", "รวมค่าโดยสาร", "ระยะทาง", "จุดรับ", "จุดส่ง",
        "รูปส่งลูกค้า", "ไฟล์", "ลิงก์รูป"]


def n(v):
    return v or 0


def line(r):
    return [r["id"], r.get("driver_name"), r.get("category"), r.get("service_type"), r.get("trip_date"),
            r.get("trip_time"), r.get("booking_code"), r.get("base_fare"), r.get("bonus"), r.get("turbo"),
            r.get("tip"), r.get("tolls"), r.get("net_earnings"), r.get("passenger_paid"),
            r.get("passenger_total"), r.get("distance_km"), r.get("pickup_text"), r.get("dropoff_text"),
            r.get("customer_image"), r.get("file_name"), r.get("source_url")]


def audit(rows, pictures=None):
    """{sheet name: [(group key, [row, ...]) or row]} and the summary lines."""
    out = {}
    # 1. one booking code on two rows of the same week
    by_code = defaultdict(list)
    for r in rows:
        if (r.get("booking_code") or "").strip():
            by_code[r["booking_code"].strip().upper()].append(r)
    out["งานซ้ำ-รหัสการจอง"] = [g for g in by_code.values() if len(g) > 1]
    # 2. equal in every column the customer sees, and no code to tell them apart
    key = lambda r: (r.get("trip_date"), r.get("trip_time"), r.get("base_fare"), r.get("distance_km"),
                     (r.get("pickup_text") or "").strip())
    same = defaultdict(list)
    for r in rows:
        if r.get("base_fare") and r.get("distance_km"):
            same[key(r)].append(r)
    twins, not_dup = [], []
    for g in same.values():
        if len(g) < 2:
            continue
        codes = [(r.get("booking_code") or "").strip().upper() for r in g]
        if all(codes) and len(set(codes)) == len(codes):
            not_dup.append(g)            # two real trips that happen to look alike
        elif all(codes):
            continue                     # the code repeats: already on the booking-code sheet
        else:
            twins.append(g)              # a row with no code — only a person can tell
    out["เหมือนทุกช่อง-ไม่มีรหัสยืนยัน"] = twins
    out["เหมือนทุกช่อง-แต่รหัสต่างกัน"] = not_dup
    # 3. money
    out["ค่ารอบเป็น0"] = [r for r in rows if not r.get("base_fare")]
    out["คุณได้รับไม่ตรงค่ารอบ+โบนัส"] = [
        r for r in rows if r.get("base_fare") and r.get("net_earnings") is not None
        and abs(n(r["net_earnings"]) - (n(r["base_fare"]) + n(r["bonus"]) + n(r["turbo"]))) > 0.01
        and abs(n(r["net_earnings"]) - (n(r["base_fare"]) + n(r["bonus"]) + n(r["turbo"]) + n(r["tolls"]))) > 0.01]
    out["บล็อกผู้โดยสารไม่ลงตัว"] = [r for r in rows if (g := gap(r)) is not None and abs(g) >= 0.51]
    # 4. what the file cannot show
    out["ข้อมูลไม่ครบ"] = [r for r in rows if not (r.get("pickup_text") and r.get("dropoff_text"))
                          or not r.get("distance_km") or not r.get("trip_time") or not r.get("service_type")]
    out["กลุ่มไม่ตรงServiceType"] = [r for r in rows if GROUP_OF.get((r.get("service_type") or "").strip())
                                     not in (None, (r.get("category") or "").strip())]
    # 5. pictures
    names = Counter(r.get("customer_image") for r in rows if r.get("customer_image"))
    cat_names = Counter((r.get("category"), r.get("customer_image")) for r in rows if r.get("customer_image"))
    out["ไม่มีรูปส่งลูกค้า"] = [r for r in rows if not r.get("customer_image")]
    out["ชื่อรูปซ้ำในกลุ่มเดียวกัน"] = [r for r in rows if cat_names[(r.get("category"), r.get("customer_image"))] > 1]
    out["ชื่อรูปซ้ำข้ามกลุ่ม"] = [r for r in rows if r.get("customer_image") and names[r["customer_image"]] > 1
                                 and cat_names[(r.get("category"), r["customer_image"])] == 1]
    if pictures is not None:
        out["รูปหายจาก Drive"] = [r for r in rows if r.get("customer_image")
                                  and r["customer_image"] not in pictures.get(r.get("category"), set())]
        named = {(r.get("category"), r.get("customer_image")) for r in rows}
        out["รูปใน Drive ที่ไม่มีแถว"] = sorted((g, p) for g, ps in pictures.items() for p in ps if (g, p) not in named)
    return out


def pictures_on_drive(drive, exports_id, d_from):
    wk = next((f["id"] for f in drive.list_folders(exports_id) if f["name"] == week_label(d_from)), None)
    if not wk:
        return None
    return {g["name"].strip(): {i["name"] for i in drive.list_images(g["id"])} for g in drive.list_folders(wk)}


def summary(rows, out):
    groups = lambda v: v and isinstance(v[0], list)
    lines = [f"{len(rows):,} แถว · " + " · ".join(f"{k} {v:,}" for k, v in Counter(r.get("service_type") for r in rows).most_common())]
    for name, v in out.items():
        if groups(v):
            extra = sum(sum(n(r.get("base_fare")) for r in g[1:]) for g in v)
            lines.append(f"  {name:<34} {len(v):>5} ชุด · {sum(len(g) for g in v):,} แถว · ค่ารอบของใบที่เกิน ฿{extra:,.0f}")
        else:
            lines.append(f"  {name:<34} {len(v):>5}")
    return lines


def build_xlsx(rows, out, d_from):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "สรุป"
    for s in [f"ตรวจ {week_label(d_from)}"] + summary(rows, out):
        ws.append([s])
    head, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E3A5F")
    for name, v in out.items():
        sh = wb.create_sheet(name[:31])
        if name == "รูปใน Drive ที่ไม่มีแถว":
            sh.append(["กลุ่ม", "ชื่อรูป"])
            for g, p in v:
                sh.append([g, p])
        else:
            grouped = bool(v) and isinstance(v[0], list)
            sh.append((["ชุด"] if grouped else []) + COLS)
            for i, g in enumerate(v if grouped else [[r] for r in v], 1):
                for r in g:
                    sh.append(([i] if grouped else []) + line(r))
        for c in sh[1]:
            c.font, c.fill = head, fill
        sh.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def main(argv=None):
    ap = argparse.ArgumentParser(description="ตรวจสัปดาห์ที่ส่งแล้วจากฐานข้อมูล (อ่านอย่างเดียว)")
    ap.add_argument("--from", dest="d_from", required=True)
    ap.add_argument("--to", dest="d_to", required=True)
    ap.add_argument("--to-drive", action="store_true")
    ap.add_argument("--xlsx")
    a = ap.parse_args(argv)
    sys.stdout.reconfigure(encoding="utf-8")
    db.init_db()
    rows = week_rows(a.d_from, a.d_to)
    drive = pics = None
    if a.to_drive:
        import config
        import roster
        drive = roster._drive()
        pics = pictures_on_drive(drive, config.DRIVE_EXPORTS_FOLDER_ID, a.d_from)
    out = audit(rows, pics)
    print(f"ตรวจ {week_label(a.d_from)} ({a.d_from}..{a.d_to})")
    print("\n".join(summary(rows, out)))
    for g in out["งานซ้ำ-รหัสการจอง"][:15]:
        print("    รหัสซ้ำ", g[0].get("booking_code"), " | ".join(f"#{r['id']} {r.get('driver_name')} ฿{r.get('base_fare')}" for r in g))
    for g in out["เหมือนทุกช่อง-ไม่มีรหัสยืนยัน"][:10]:
        print("    เหมือนกัน", " | ".join(f"#{r['id']} {r.get('driver_name')} ฿{r.get('base_fare')} {r.get('booking_code') or '-'}" for r in g))
    data = build_xlsx(rows, out, a.d_from)
    if a.xlsx:
        open(a.xlsx, "wb").write(data)
    if drive:
        import config
        folder = drive.ensure_folder(config.DRIVE_EXPORTS_FOLDER_ID, DRIVE_DIR)
        fid = drive.upload_xlsx(folder, f"ตรวจ {week_label(a.d_from)}.xlsx", data)
        print(f"\nวางรายงานลง Drive: {DRIVE_DIR}/ตรวจ {week_label(a.d_from)}.xlsx → https://drive.google.com/file/d/{fid}/view")
    return 0


if __name__ == "__main__":
    sys.exit(main())
