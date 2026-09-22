# -*- coding: utf-8 -*-
"""The week audit sorts a delivered week's faults onto their own lists, and changes nothing."""
import io
import os
import shutil
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
WORK = tempfile.mkdtemp(prefix="pocr-audit-")
os.environ["PHOTO_OCR_DATA"] = WORK
os.environ.pop("DATABASE_URL", None)
sys.path.insert(0, "backend")

from openpyxl import load_workbook                              # noqa: E402
from sqlalchemy import insert, select                           # noqa: E402

import db                                                       # noqa: E402
import week_audit as wa                                         # noqa: E402

db.init_db()
ok = True


def check(name, cond):
    global ok
    print(("✅" if cond else "✗"), name)
    ok = ok and cond


W = ("2026-08-17", "2026-08-23")
a = db.create_job("พงศ์กฤษณ์", "Trips", *W, category="2 W Standard")
b = db.create_job("วิมลวรรณ", "Trips", *W, category="2 W Standard")
s = db.create_job("ณัฐพล", "Trips", *W, category="2 W Saver")
c = db.create_job("ณัฐพล", "Trips", *W, category="4 W Standard")
base = dict(status="done", committed=1, file_name="x.jpg", trip_date="2026-08-17", trip_time="12:00-17:59",
            service_type="Standard Bike", distance_km=15.7, pickup_text="A", dropoff_text="B",
            passenger_paid=140.0, passenger_total=130.0, app_fee=-1.0, discount=None, other_adj=-5.0)
with db.engine.begin() as x:
    ins = lambda **kw: x.execute(insert(db.trips).values(**{**base, **kw})).inserted_primary_key[0]
    ins(job_id=a, base_fare=118.0, net_earnings=118.0, booking_code="A-1", customer_image="WK34-พงศ์กฤษณ์1.jpg")
    ins(job_id=b, base_fare=118.0, net_earnings=118.0, booking_code="a-1 ", customer_image="WK34-วิมลวรรณ1.jpg")
    ins(job_id=a, base_fare=39.0, net_earnings=39.0, distance_km=5.29, booking_code=None, customer_image="WK34-พงศ์กฤษณ์2.jpg")
    ins(job_id=b, base_fare=39.0, net_earnings=39.0, distance_km=5.29, booking_code="A-2", customer_image="WK34-วิมลวรรณ2.jpg")
    ins(job_id=a, base_fare=26.0, net_earnings=26.0, distance_km=1.3, booking_code="A-3", customer_image="WK34-พงศ์กฤษณ์3.jpg")
    ins(job_id=b, base_fare=26.0, net_earnings=26.0, distance_km=1.3, booking_code="A-4", customer_image="WK34-วิมลวรรณ3.jpg")
    zero = ins(job_id=b, base_fare=0.0, net_earnings=5.0, turbo=5.0, distance_km=0, booking_code="A-5",
               customer_image="WK34-วิมลวรรณ4.jpg", passenger_paid=None)
    ins(job_id=s, base_fare=30.0, net_earnings=30.0, service_type="Saver Bike", distance_km=2.0, booking_code="A-6",
        customer_image="WK34-ณัฐพล1.jpg")
    ins(job_id=c, base_fare=90.0, net_earnings=99.0, service_type="Standard Car", distance_km=9.0, booking_code="A-7",
        customer_image="WK34-ณัฐพล1.jpg")
    ins(job_id=c, base_fare=50.0, net_earnings=50.0, service_type="Saver Bike", distance_km=4.0, booking_code="A-8",
        customer_image=None)

rows = wa.week_rows(*W)
out = wa.audit(rows, pictures={"2 W Standard": {"WK34-พงศ์กฤษณ์1.jpg", "WK34-เกิน.jpg"}})
check("❗ รหัสการจองเดียวกัน (ต่างตัวพิมพ์/ช่องว่าง) คือซ้ำ", len(out["งานซ้ำ-รหัสการจอง"]) == 1
      and len(out["งานซ้ำ-รหัสการจอง"][0]) == 2)
check("เหมือนทุกช่องแต่แถวหนึ่งไม่มีรหัส → ให้คนดู", len(out["เหมือนทุกช่อง-ไม่มีรหัสยืนยัน"]) == 1)
check("เหมือนทุกช่องแต่รหัสต่างกัน → งานจริงสองงาน แยกไว้", len(out["เหมือนทุกช่อง-แต่รหัสต่างกัน"]) == 1)
check("ค่ารอบ 0", [r["id"] for r in out["ค่ารอบเป็น0"]] == [zero])
check("คุณได้รับไม่ตรงค่ารอบ+โบนัส", len(out["คุณได้รับไม่ตรงค่ารอบ+โบนัส"]) == 1)
check("บล็อกผู้โดยสารไม่ลงตัว (130 ≠ 140 − 1 − 5)", len(out["บล็อกผู้โดยสารไม่ลงตัว"]) >= 1)
check("กลุ่มไม่ตรง Service Type", len(out["กลุ่มไม่ตรงServiceType"]) == 1)
check("ชื่อรูปเดียวกันคนละกลุ่ม แยกจากชื่อซ้ำในกลุ่มเดียวกัน",
      len(out["ชื่อรูปซ้ำข้ามกลุ่ม"]) == 2 and not out["ชื่อรูปซ้ำในกลุ่มเดียวกัน"])
check("ไม่มีรูปส่งลูกค้า", len(out["ไม่มีรูปส่งลูกค้า"]) == 1)
check("รูปหายจาก Drive และรูปที่ไม่มีแถว", out["รูปใน Drive ที่ไม่มีแถว"] == [("2 W Standard", "WK34-เกิน.jpg")]
      and len(out["รูปหายจาก Drive"]) == 8)
wb = load_workbook(io.BytesIO(wa.build_xlsx(rows, out, W[0])))
check("Excel มีชีตสรุปและชีตรายการ", wb.sheetnames[0] == "สรุป" and "งานซ้ำ-รหัสการจอง" in wb.sheetnames)
with db.engine.begin() as x:
    before = x.execute(select(db.trips)).mappings().all()
wa.main(["--from", W[0], "--to", W[1]])
with db.engine.begin() as x:
    after = x.execute(select(db.trips)).mappings().all()
check("อ่านอย่างเดียว ไม่แก้แถวใด", [dict(r) for r in before] == [dict(r) for r in after])

shutil.rmtree(WORK, ignore_errors=True)
print("\nสรุป:", "ผ่านทั้งหมด ✅" if ok else "มีข้อที่ไม่ผ่าน ✗")
sys.exit(0 if ok else 1)
