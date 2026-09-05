# -*- coding: utf-8 -*-
"""Two requirements from Ops on 2026-09-05.

1. Repeated work comes out. Until now only the same rider repeating a booking code inside the
   same week counted; the same slip handed in under two rider names flowed through, and over
   W33-W35 that was 1,521 rows (11.6%). Forward-only by Ops' decision — delivered weeks are
   left alone, so these tests are about what happens to work arriving from now on.
2. Sheet1's pick-up and drop-off show the place from the slip, not the zone (which said
   "Downtown" for 10,942 of 13,077 rows).
"""
import os
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
WORK = tempfile.mkdtemp(prefix="pocr-ops-")
os.environ["PHOTO_OCR_DATA"] = WORK
os.environ.pop("DATABASE_URL", None)
sys.path.insert(0, "backend")

import config                                                   # noqa: E402
import db                                                       # noqa: E402
import excel_writer                                             # noqa: E402

db.init_db()
ok = True


def check(name, cond):
    global ok
    print(("✅" if cond else "✗"), name)
    ok = ok and cond


def trip(job, code, base=100, when="14:30", date="2026-09-01", name=None):
    t = db.create_trip(job, name or f"{code}.jpg", None, "image/jpeg")
    db.update_trip(t, {"status": "done", "trip_date": date, "trip_time": when,
                       "net_earnings": base, "base_fare": base, "booking_code": code,
                       "check_status": "pass"})
    return t


WEEK = ("2026-08-31", "2026-09-06")
OTHER = ("2026-09-07", "2026-09-13")

# --- 1. งานซ้ำ -------------------------------------------------------------------------------
a = db.create_job("สมชาย", "Trips", *WEEK)
t1 = trip(a, "A-9AAAAAAAAAAAAAV")
db.auto_approve_job(a)
check("ใบแรกอนุมัติผ่านปกติ", db.get_trip(t1)["committed"] == 1)

# คนละไรเดอร์ สัปดาห์เดียวกัน — เดิมปล่อยผ่าน ตอนนี้ต้องถูกตัด
b = db.create_job("สมหญิง", "Trips", *WEEK)
t2 = trip(b, "A-9AAAAAAAAAAAAAV")
db.auto_approve_job(b)
r2 = db.get_trip(t2)
check("คนละไรเดอร์ code เดียวกัน: ไม่อนุมัติ", r2["committed"] == 0)
check("คนละไรเดอร์ code เดียวกัน: ถูกทิ้งออกจากคิว", r2["status"] == "duplicate")
check("บอกไว้ว่าซ้ำกับใบไหน", "ซ้ำกับ" in (r2["note"] or ""))
check("ใบแรกยังอยู่ครบ (เก็บของคนแรก)", db.get_trip(t1)["committed"] == 1)

# คนเดียวกัน คนละสัปดาห์ — เดิม Ops อนุญาต ตอนนี้ต้องถูกตัด
c = db.create_job("สมชาย", "Trips", *OTHER)
t3 = trip(c, "A-9AAAAAAAAAAAAAV", date="2026-09-08")
db.auto_approve_job(c)
check("คนเดียวกัน คนละสัปดาห์: ไม่อนุมัติ", db.get_trip(t3)["committed"] == 0)

# code เดียวกันแต่ยอดไม่ตรง = คำถามจริง ต้องให้คนดู ไม่ใช่ทิ้งเงียบ
d = db.create_job("สมศรี", "Trips", *WEEK)
t4 = trip(d, "A-9AAAAAAAAAAAAAV", base=250)
db.auto_approve_job(d)
r4 = db.get_trip(t4)
check("code ซ้ำแต่ยอดต่างกัน: ไม่อนุมัติ", r4["committed"] == 0)
check("code ซ้ำแต่ยอดต่างกัน: ค้างในคิวให้คนตัดสิน ไม่ทิ้งเอง", r4["status"] == "done")

# code สั้น (ตกบรรทัด) เป็นแค่เบาะแส ต้องมียอด+เวลาตรงด้วยถึงจะนับว่าซ้ำ
e = db.create_job("วิชัย", "Trips", *WEEK)
t5 = trip(e, "A-9SHORT", base=80, when="09:15")
db.auto_approve_job(e)
check("code สั้นใบแรก: อนุมัติปกติ", db.get_trip(t5)["committed"] == 1)

f = db.create_job("วิภา", "Trips", *WEEK)
t6 = trip(f, "A-9SHORT", base=80, when="09:15")
t7 = trip(f, "A-9SHORT", base=145, when="18:40")      # code สั้นชนกัน แต่คนละทริปจริง
db.auto_approve_job(f)
check("code สั้น + ยอด/เวลาตรง: นับว่าซ้ำ", db.get_trip(t6)["committed"] == 0)
check("code สั้น + ยอด/เวลาต่าง: ไม่ตัดผิด อนุมัติได้", db.get_trip(t7)["committed"] == 1)

# ไม่มี code เลย — จับด้วยวิธีนี้ไม่ได้ ต้องไม่พังและต้องไม่บล็อกมั่ว
g = db.create_job("อนันต์", "Trips", *WEEK)
t8 = db.create_trip(g, "nocode.jpg", None, "image/jpeg")
db.update_trip(t8, {"status": "done", "trip_date": "2026-09-02", "net_earnings": 60,
                    "base_fare": 60, "check_status": "pass"})
db.auto_approve_job(g)
check("ไม่มี booking code: ยังอนุมัติได้ตามปกติ", db.get_trip(t8)["committed"] == 1)

check("ใบที่ถูกทิ้งโผล่ในบันทึกให้กู้คืนได้",
      any(x["id"] == t2 for x in db.discarded_duplicates()))
db.restore_discarded(t2)
check("กู้กลับเข้าคิวได้", db.get_trip(t2)["status"] == "done")

# --- 2. สถานที่บน Sheet1 ---------------------------------------------------------------------
check("ไม่ได้ปิดการอ่านที่อยู่แล้ว", "pickup_text" not in config.EXTRACT_DROP_FIELDS)
check("มีที่อยู่: ใช้ที่อยู่", excel_writer.place("เอ็มสเฟียร์ ทางออกป้ายรถเมล์", "วัฒนา", "BKK")
      == "เอ็มสเฟียร์ ทางออกป้ายรถเมล์")
check("ไม่มีที่อยู่: ถอยไปใช้โซนเดิม ไม่ปล่อยว่าง",
      excel_writer.place(None, "ดอนเมือง", "BKK") == "North-DMK")
check("ที่อยู่เป็นช่องว่างล้วน: ถือว่าไม่มี", excel_writer.place("   ", "ดอนเมือง", "BKK") == "North-DMK")

check("Analysis ไม่มีคอลัมน์ที่อยู่แล้ว",
      "Pick-up Address" not in excel_writer.ANALYSIS_HEADERS
      and "Drop-off Address" not in excel_writer.ANALYSIS_HEADERS)
check("Analysis ยังเก็บโซนไว้", "Pick-up Zone" in excel_writer.ANALYSIS_HEADERS)
check("หัวตาราง Analysis กับความกว้างยังเท่ากัน",
      len(excel_writer.ANALYSIS_HEADERS) == len(excel_writer.ANALYSIS_WIDTHS))

rows = [{"driver_name": "สมชาย", "trip_date": "2026-09-01", "trip_time": "14:30",
         "service_type": "Standard Bike", "payment_method": "CASH", "base_fare": 100,
         "net_earnings": 100, "distance_km": 5.0, "passenger_total": 117,
         "pickup_text": "เอ็มสเฟียร์ ทางออกป้ายรถเมล์", "dropoff_text": "มหานคร สกายวอล์ค",
         "pickup_district": "วัฒนา", "dropoff_district": "บางรัก", "booking_code": "A-9ZZZ"}]
import io                                                      # noqa: E402
import openpyxl                                                 # noqa: E402
wb = openpyxl.load_workbook(io.BytesIO(excel_writer.build_workbook(rows)))
ws, wa = wb[excel_writer.SHEET], wb[excel_writer.ANALYSIS_SHEET]
check("Sheet1 ต้นทางเป็นสถานที่จริง", ws.cell(2, 6).value == "เอ็มสเฟียร์ ทางออกป้ายรถเมล์")
check("Sheet1 ปลายทางเป็นสถานที่จริง", ws.cell(2, 7).value == "มหานคร สกายวอล์ค")
check("จำนวนคอลัมน์ Analysis ตรงกับหัวตาราง",
      sum(1 for i in range(1, len(excel_writer.ANALYSIS_HEADERS) + 2)
          if wa.cell(1, i).value) == len(excel_writer.ANALYSIS_HEADERS))
check("แถว Analysis ไม่มีที่อยู่หลุดมา",
      all("เอ็มสเฟียร์" not in str(wa.cell(2, i).value or "")
          for i in range(1, len(excel_writer.ANALYSIS_HEADERS) + 1)))

print("\nสรุป:", "ผ่านทั้งหมด ✅" if ok else "มีข้อที่ไม่ผ่าน ✗")
sys.exit(0 if ok else 1)
