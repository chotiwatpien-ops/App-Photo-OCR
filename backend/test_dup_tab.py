# -*- coding: utf-8 -*-
"""The 'งานซ้ำ' tab's detail panel and report: each repeat beside the picture it repeats, the
read ones can be put back, and the workbook downloads straight from the page (redesign 2026-10)."""
import io
import os
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
os.environ["PHOTO_OCR_DATA"] = tempfile.mkdtemp(prefix="pocr-duptab-")
os.environ.pop("DATABASE_URL", None)
sys.path.insert(0, "backend")

from sqlalchemy import insert                                   # noqa: E402

import config                                                   # noqa: E402
import db                                                       # noqa: E402
import main                                                     # noqa: E402
from fastapi.testclient import TestClient                       # noqa: E402

db.init_db()
config.APP_PASSWORD = ""
ok = True


def check(name, cond):
    global ok
    print(("✅" if cond else "✗"), name)
    ok = ok and cond


W = ("2026-09-28", "2026-10-04")
job = db.create_job("มานพ จ.", "Trips", *W, category="4 W Standard")


def trip(**kw):
    with db.engine.begin() as c:
        return c.execute(insert(db.trips).values(job_id=job, **kw)).inserted_primary_key[0]


first = trip(file_name="4W-Taxi Meen ชุด 1_1+2_฿80.jpg", status="done", committed=1, booking_code="A-5KX19P",
             source_album="4W-Taxi Meen ชุด 1", net_earnings=80, trip_date="2026-09-29")
again = trip(file_name="4W-Taxi Meen ชุด 2_5+6_฿80.jpg", status="duplicate", committed=0, booking_code="A-5KX19P",
             source_album="4W-Taxi Meen ชุด 2", net_earnings=80, trip_date="2026-09-29", duplicate_of=first)
db.record_skipped_copies([{"ref": "pool:x/4W-Taxi Meen ชุด 2/a.jpg", "week_from": W[0], "album": "4W-Taxi Meen ชุด 2",
                           "file_name": "a.jpg", "drive_id": "D1", "same_as": "4W-Taxi Meen ชุด 1/a.jpg",
                           "same_as_album": "4W-Taxi Meen ชุด 1", "kind": "สำเนาจากอัลบั้มอื่น"}])

web = TestClient(main.app)
albums = web.get(f"/api/duplicates/albums?week={W[0]}").json()["albums"]
a2 = next(a for a in albums if a["album"] == "4W-Taxi Meen ชุด 2")
check("อัลบั้มบอกซ้ำก่อนอ่าน 1 · หลังอ่าน 1", (a2["pre"], a2["post"], a2["dup"]) == (1, 1, 2))

det = web.get("/api/duplicates/album", params={"week": W[0], "album": "4W-Taxi Meen ชุด 2"}).json()
read = [p for p in det["pairs"] if not p.get("pre")]
pre = [p for p in det["pairs"] if p.get("pre")]
check("รายละเอียด: ใบที่อ่านแล้ว 1 · ไฟล์เดิมส่งซ้ำ 1", len(read) == 1 and len(pre) == 1)
check("❗ ใบที่อ่านแล้วรู้ว่าซ้ำกับแถวไหน (ไว้โชว์รูปคู่กัน)", read[0]["trip_id"] == again and read[0]["twin_id"] == first)
check("บอกเลขจองและอัลบั้มของใบเดิม", read[0]["code"] == "A-5KX19P" and read[0]["twin_album"] == "4W-Taxi Meen ชุด 1")
check("ใบที่อ่านแล้วกดเก็บไว้ได้ · ไฟล์เดิมที่ไม่ได้อ่านกดไม่ได้", read[0]["restorable"] and not pre[0]["restorable"])
check("ไฟล์เดิมส่งซ้ำมีลิงก์ไป Drive", pre[0]["link"] == "https://drive.google.com/file/d/D1/view")

res = web.get(f"/api/duplicates/report.xlsx?week={W[0]}")
from openpyxl import load_workbook                              # noqa: E402
wb = load_workbook(io.BytesIO(res.content))
check("ดาวน์โหลดรายงาน Excel ได้ ชีทแรกเป็นรายอัลบั้ม", res.status_code == 200 and wb.sheetnames[0] == "รายอัลบั้ม")
check("ชื่อไฟล์เป็นภาษาคน (งานซ้ำ 2026 WK40)", "WK40" in res.headers["content-disposition"])
check("❗ คอลัมน์ภายใน twin_id ไม่หลุดไปในไฟล์", all("twin_id" not in [c.value for c in ws[1]] for ws in wb.worksheets))

r = web.post(f"/api/trips/{again}/restore")
a2 = next(a for a in web.get(f"/api/duplicates/albums?week={W[0]}").json()["albums"] if a["album"] == "4W-Taxi Meen ชุด 2")
check("❗ กดเก็บไว้แล้ว ตัวเลขซ้ำอัปเดตทันที ไม่ต้องรอแคช", r.status_code == 200 and a2["post"] == 0)
check("สัปดาห์ที่ไม่มีอยู่ → 404", web.get("/api/duplicates/album", params={"week": "2020-01-06", "album": "x"}).status_code == 404)

print("\nผ่านทั้งหมด" if ok else "\n✗ มีข้อที่ไม่ผ่าน")
sys.exit(0 if ok else 1)
