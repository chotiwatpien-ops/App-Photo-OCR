# -*- coding: utf-8 -*-
"""รายงานรูปซ้ำรายสัปดาห์ — ให้ลูกค้าตีกลับไรเดอร์ได้ (2026-09-12, ระยะ 1)

อ่านอย่างเดียว ไม่แก้ไม่ลบอะไรทั้งสิ้น

ระบบพักรูปซ้ำไว้เป็น status='duplicate' พร้อมโน้ตว่าไปซ้ำกับไฟล์ไหน แต่โน้ตเป็นข้อความ ค้นไม่ได้
และไม่มีที่บันทึกว่าตีกลับไปหรือยัง ตัวนี้ประกอบเคสให้ครบ: ใบไหน ของไรเดอร์คนไหน ซ้ำกับงานไหน
ของใคร สัปดาห์ไหน แล้วเขียนเป็น Excel ให้ส่งต่อได้

กติกาที่ Ops ตอบไว้ 2026-09-12: ตีกลับ 'ไรเดอร์เจ้าของอัลบั้มที่ส่งมา' และเอา 'คนที่ส่งทีหลัง'
ระบบอนุมัติใบแรกที่เข้ามาเสมอ ใบที่ถูกพักจึงเป็นของคนที่ส่งทีหลังอยู่แล้วโดยโครงสร้าง

ระยะถัดไป (ยังไม่ได้ทำ): ตาราง dup_cases เก็บสถานะการตีกลับและคำตอบผู้ส่ง + ปุ่มในหน้า Admin
แล้วรายงานนี้จะอ่านสถานะจากตารางนั้นแทนคอลัมน์ว่าง

    python duplicate_report.py --from 2026-09-07 --to 2026-09-13
    python duplicate_report.py --from 2026-09-07 --to 2026-09-13 --xlsx dup_w37.xlsx
"""
import argparse
import collections
import re
import sys

from sqlalchemy import select

import db

# 'LINE_ALBUM_2W-Win ป๋อง = 71_86659+86658_฿30.jpg' -> ชื่ออัลบั้มที่ไรเดอร์ส่งมา
ALBUM = re.compile(r"^(?:LINE_ALBUM_)?(.+?)_\d+\+\d+_฿")
KIND_BOTTOM = "ครึ่งล่างของงานที่นับไปแล้ว"
KIND_SAME_WEEK = "รหัสการจองซ้ำในสัปดาห์เดียวกัน"
KIND_OTHER_WEEK = "รหัสการจองซ้ำข้ามสัปดาห์"
KIND_SAME_FILE = "ไฟล์เดิมส่งซ้ำ"
KINDS = (KIND_SAME_FILE, KIND_SAME_WEEK, KIND_OTHER_WEEK, KIND_BOTTOM)


def album_of(file_name, folder_name=None, driver_name=None):
    """ไรเดอร์เจ้าของอัลบั้มที่ส่งรูปใบนี้มา

    รูปที่ pool ต่อให้พกชื่ออัลบั้มไว้ในชื่อไฟล์ ซึ่งคือ 'คนส่ง' ตัวจริง — ต่างจากโฟลเดอร์ที่ระบบ
    เอารูปไปวาง เพราะระบบกระจายงานให้ไรเดอร์คนไหนก็ได้ที่ยังมีที่ว่าง (รอบ W37 อัลบั้มของ ป๋อง
    ไปโผล่ในโฟลเดอร์ของสรธรและธีรพงศ์) รูปเดี่ยวไม่มีชื่ออัลบั้มติดมา จึงได้แค่โฟลเดอร์ต้นทาง"""
    m = ALBUM.match(file_name or "")
    if m:
        return m.group(1).strip()
    return (folder_name or driver_name or "").strip() or "ไม่รู้อัลบั้ม"


def kind_of(row, twin):
    if (row.get("note") or "").startswith("ครึ่งล่างของงานที่อนุมัติแล้ว"):
        return KIND_BOTTOM
    if twin and row.get("image_hash") and row["image_hash"] == twin.get("image_hash"):
        return KIND_SAME_FILE
    if twin and twin.get("date_from") and row.get("date_from") != twin["date_from"]:
        return KIND_OTHER_WEEK
    return KIND_SAME_WEEK


COLS = ["id", "ไฟล์ที่ซ้ำ", "ไรเดอร์ผู้ส่ง", "โฟลเดอร์ที่ระบบวาง", "กลุ่ม", "สัปดาห์",
        "รหัสการจอง", "ยอด", "วันที่งาน", "ซ้ำกับไฟล์", "ของไรเดอร์", "สัปดาห์ของต้นฉบับ",
        "รูปที่ส่งลูกค้า", "ประเภทการซ้ำ", "ลิงก์รูป", "สถานะเคส", "ผู้ตีกลับ",
        "วันที่ตีกลับ", "คำตอบผู้ส่ง"]


def build(rows, twins):
    """(เคสที่พร้อมตีกลับ, เคสที่ต้องให้คนตัดสิน) — twins: id -> แถวต้นฉบับที่อนุมัติแล้ว"""
    cases, undecided = [], []
    for r in rows:
        twin = twins.get(r.get("duplicate_of")) or twins.get(r.get("twin_id"))
        t = twin or {}
        case = {
            "id": r["id"],
            "ไฟล์ที่ซ้ำ": r["file_name"],
            "ไรเดอร์ผู้ส่ง": album_of(r["file_name"], r.get("folder_name"), r.get("driver_name")),
            "โฟลเดอร์ที่ระบบวาง": r.get("driver_name"),
            "กลุ่ม": r.get("category"),
            "สัปดาห์": r.get("date_from"),
            "รหัสการจอง": r.get("booking_code"),
            "ยอด": r.get("net_earnings"),
            "วันที่งาน": r.get("trip_date"),
            "ซ้ำกับไฟล์": t.get("file_name"),
            "ของไรเดอร์": album_of(t.get("file_name"), None, t.get("driver_name")) if twin else None,
            "สัปดาห์ของต้นฉบับ": t.get("date_from"),
            "รูปที่ส่งลูกค้า": t.get("customer_image"),
            "ประเภทการซ้ำ": kind_of(r, twin),
            "ลิงก์รูป": r.get("source_url"),
            "สถานะเคส": "รอตีกลับ",
            "ผู้ตีกลับ": "", "วันที่ตีกลับ": "", "คำตอบผู้ส่ง": "",
        }
        (cases if r["status"] == "duplicate" else undecided).append(case)
    return cases, undecided


def summary(cases):
    """สรุปรายไรเดอร์ผู้ส่ง — แถวที่ลูกค้าใช้คุยกับคนส่ง"""
    by = collections.defaultdict(collections.Counter)
    for c in cases:
        s = by[c["ไรเดอร์ผู้ส่ง"]]
        s["ใบซ้ำ"] += 1
        s["ยอดรวม"] += c["ยอด"] or 0
        s[c["ประเภทการซ้ำ"]] += 1
        if c["ของไรเดอร์"] and c["ของไรเดอร์"] != c["ไรเดอร์ผู้ส่ง"]:
            s["ซ้ำใต้ชื่อคนอื่น"] += 1
    return by


def load(d_from, d_to):
    """แถวซ้ำของสัปดาห์ + ต้นฉบับของมัน (ต้นฉบับอยู่สัปดาห์ไหนก็ได้)"""
    t, j = db.trips.c, db.jobs.c
    cols = [t.id, t.file_name, t.status, t.committed, t.booking_code, t.duplicate_of,
            t.net_earnings, t.trip_date, t.note, t.customer_image, t.source_url, t.image_hash,
            j.driver_name, j.category, j.folder_name, j.date_from]
    src = db.trips.join(db.jobs, j.id == t.job_id)
    with db.engine.begin() as c:
        week = [dict(r) for r in c.execute(
            select(*cols).select_from(src)
            .where(j.date_from >= d_from, j.date_from <= d_to,
                   (t.status == "duplicate") | (t.duplicate_of.isnot(None)))).mappings().all()]
        ids = [r["duplicate_of"] for r in week if r["duplicate_of"]]
        codes = [r["booking_code"] for r in week if r["booking_code"]]
        twins = [dict(r) for r in c.execute(
            select(*cols).select_from(src).where(t.id.in_(ids))).mappings().all()] if ids else []
        by_code = [dict(r) for r in c.execute(
            select(*cols).select_from(src)
            .where(t.committed == 1, t.booking_code.in_(codes))).mappings().all()] if codes else []
    twin_by_id = {r["id"]: r for r in twins}
    code_twin = {}
    for r in sorted(by_code, key=lambda r: r["id"]):     # ใบแรกที่อนุมัติคือต้นฉบับ
        code_twin.setdefault(r["booking_code"], r)
    for r in week:                                      # เติมต้นฉบับให้แถวที่ไม่มี duplicate_of
        if not r["duplicate_of"] and r["booking_code"]:
            tw = code_twin.get(r["booking_code"])
            if tw and tw["id"] != r["id"]:
                r["twin_id"] = tw["id"]
                twin_by_id[tw["id"]] = tw
    return week, twin_by_id


def write_xlsx(path, cases, undecided):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    head, fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="1E3A5F")

    def sheet(ws, headers, rows_):
        ws.append(headers)
        for cell in ws[1]:
            cell.font, cell.fill = head, fill
            cell.alignment = Alignment(vertical="center")
        for r in rows_:
            ws.append(r)
        for i, h in enumerate(headers, 1):
            ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(46, len(str(h)) + 8))
        ws.freeze_panes = "A2"
        return ws

    sheet(wb.active, ["ไรเดอร์ผู้ส่ง", "ใบซ้ำ", "ยอดรวม", KIND_SAME_FILE, KIND_SAME_WEEK,
                      KIND_OTHER_WEEK, KIND_BOTTOM, "ซ้ำใต้ชื่อคนอื่น"],
          [[k, v["ใบซ้ำ"], round(v["ยอดรวม"], 2), v[KIND_SAME_FILE], v[KIND_SAME_WEEK],
            v[KIND_OTHER_WEEK], v[KIND_BOTTOM], v["ซ้ำใต้ชื่อคนอื่น"]]
           for k, v in sorted(summary(cases).items(), key=lambda kv: -kv[1]["ใบซ้ำ"])])
    wb.active.title = "สรุปรายผู้ส่ง"
    sheet(wb.create_sheet("รายใบ"), COLS, [[c.get(k) for k in COLS] for c in cases])
    sheet(wb.create_sheet("ค้างตัดสิน"), COLS, [[c.get(k) for k in COLS] for c in undecided])
    wb.save(path)
    return path


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument("--from", dest="d_from", required=True)
    ap.add_argument("--to", dest="d_to", required=True)
    ap.add_argument("--xlsx", help="เขียนไฟล์ Excel ไว้ที่นี่")
    a = ap.parse_args(argv)

    db.init_db()
    rows, twins = load(a.d_from, a.d_to)
    cases, undecided = build(rows, twins)
    print(f"สัปดาห์ {a.d_from} ถึง {a.d_to}")
    print(f"ใบซ้ำที่ระบบพักไว้ {len(cases):,} · ยังต้องให้คนตัดสิน {len(undecided):,}")
    if not cases and not undecided:
        print("ไม่พบรูปซ้ำในสัปดาห์นี้")
        return 0
    if cases:
        print(f"\n{'ไรเดอร์ผู้ส่ง':<26}{'ใบซ้ำ':>7}{'ยอดรวม':>10}  ประเภทที่เจอ")
        for k, v in sorted(summary(cases).items(), key=lambda kv: -kv[1]["ใบซ้ำ"]):
            kinds = " · ".join(f"{n} {name}" for name in KINDS for n in [v[name]] if n)
            print(f"{k[:25]:<26}{v['ใบซ้ำ']:>7}{v['ยอดรวม']:>10,.0f}  {kinds}")
    shared = [c for c in cases if c["ของไรเดอร์"] and c["ของไรเดอร์"] != c["ไรเดอร์ผู้ส่ง"]]
    if shared:
        print(f"\nในนั้น {len(shared)} ใบเป็นสลิปที่ถูกส่งใต้ชื่อไรเดอร์คนละคน:")
        for c in shared[:15]:
            print(f"  {(c['ไฟล์ที่ซ้ำ'] or '')[:46]:<48} ส่งโดย {c['ไรเดอร์ผู้ส่ง'][:18]:<20}"
                  f"ซ้ำกับของ {c['ของไรเดอร์']}")
    if a.xlsx:
        print(f"\nเขียนไฟล์แล้ว: {write_xlsx(a.xlsx, cases, undecided)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
