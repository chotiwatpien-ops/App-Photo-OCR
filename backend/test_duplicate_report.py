# -*- coding: utf-8 -*-
"""รายงานรูปซ้ำ: คนส่งคือเจ้าของอัลบั้ม ไม่ใช่โฟลเดอร์ที่ระบบเอารูปไปวาง (2026-09-12).

รอบ W37 อัลบั้มของ ป๋อง ถูกอัปโหลดซ้ำ รูป 70 ใบไปลงโฟลเดอร์ของสรธรและธีรพงศ์ ถ้ารายงานบอกชื่อ
ตามโฟลเดอร์ ลูกค้าจะไปตีกลับผิดคน ฐานข้อมูลชั่วคราวเป็น SQLite
"""
import os
import shutil
import sys
import tempfile

sys.stdout.reconfigure(encoding="utf-8")
WORK = tempfile.mkdtemp(prefix="pocr-dup-")
os.environ["PHOTO_OCR_DATA"] = WORK
os.environ.pop("DATABASE_URL", None)
sys.path.insert(0, "backend")

import db                                                       # noqa: E402
import duplicate_report as rep                                  # noqa: E402

db.init_db()
ok = True


def check(name, cond):
    global ok
    print(("✅" if cond else "✗"), name)
    ok = ok and cond


def job(driver, date_from, folder=None):
    with db.engine.begin() as c:
        return c.execute(db.jobs.insert().values(
            driver_name=driver, sheet="Trips", date_from=date_from, date_to=date_from,
            status="committed", created_at="2026-09-08 10:00:00", category="2 W Saver",
            folder_name=folder).returning(db.jobs.c.id)).scalar()


def trip(job_id, name, **kw):
    values = {"job_id": job_id, "file_name": name, "status": "done", **kw}
    with db.engine.begin() as c:
        return c.execute(db.trips.insert().values(**values).returning(db.trips.c.id)).scalar()


check("คนส่งมาจากชื่ออัลบั้มในชื่อไฟล์ ไม่ใช่โฟลเดอร์ปลายทาง",
      rep.album_of("2W-Win ป๋อง = 71_86659+86658_฿30.jpg", driver_name="สรธร") == "2W-Win ป๋อง = 71")
check("ชื่อที่ LINE ใส่ LINE_ALBUM_ นำหน้า ก็ตัดออกให้",
      rep.album_of("LINE_ALBUM_2W-Win วิน = 150_1+2_฿99.jpg") == "2W-Win วิน = 150")
check("รูปเดี่ยวไม่มีอัลบั้ม ใช้โฟลเดอร์ต้นทางแทน",
      rep.album_of("S__97419329_0.jpg", folder_name="19-ธีรพงศ์ Win") == "19-ธีรพงศ์ Win")
check("ไม่มีอะไรเลยก็ไม่ล่ม", rep.album_of(None) == "ไม่รู้อัลบั้ม")
check("ต้นทางที่ pool จดไว้ ชนะทุกอย่าง",
      rep.album_of("S__97419329_0.jpg", folder_name="19-ธีรพงศ์ Win",
                   source_album="2W-Win ป๋อง = 71") == "2W-Win ป๋อง = 71")

# เทียบ 'ส่งใต้ชื่อคนอื่น' ได้ต่อเมื่อรู้ตัวคนส่งทั้งสองฝั่ง ไม่งั้นเป็นการเอาอัลบั้มไปชนโฟลเดอร์
_sure = {"ไรเดอร์ผู้ส่ง": "ป๋อง", "ของไรเดอร์": "เอก", "ผู้ส่งแน่ชัด": True, "ต้นฉบับแน่ชัด": True}
check("รู้ทั้งสองฝั่งและคนละคน → ใช่", rep.sent_by_someone_else(_sure))
check("คนเดียวกัน → ไม่ใช่", not rep.sent_by_someone_else(dict(_sure, ของไรเดอร์="ป๋อง")))
check("ฝั่งต้นฉบับรู้แค่ชื่อโฟลเดอร์ → ไม่นับ ไม่เดา",
      not rep.sent_by_someone_else(dict(_sure, ต้นฉบับแน่ชัด=False)))
check("ฝั่งผู้ส่งรู้แค่ชื่อโฟลเดอร์ → ไม่นับ",
      not rep.sent_by_someone_else(dict(_sure, ผู้ส่งแน่ชัด=False)))

check("โน้ตครึ่งล่างบอกประเภทได้",
      rep.kind_of({"note": "ครึ่งล่างของงานที่อนุมัติแล้ว (a.jpg) — รูปเกิน ลบได้"}, None) == rep.KIND_BOTTOM)
check("hash เดียวกัน = ไฟล์เดิมส่งซ้ำ",
      rep.kind_of({"image_hash": "aa"}, {"image_hash": "aa"}) == rep.KIND_SAME_FILE)
check("ต้นฉบับอยู่คนละสัปดาห์ = ซ้ำข้ามสัปดาห์",
      rep.kind_of({"date_from": "2026-09-07"}, {"date_from": "2026-08-31"}) == rep.KIND_OTHER_WEEK)
check("ต้นฉบับสัปดาห์เดียวกัน = ซ้ำในสัปดาห์",
      rep.kind_of({"date_from": "2026-09-07"}, {"date_from": "2026-09-07"}) == rep.KIND_SAME_WEEK)

# ป๋องส่งอัลบั้มซ้ำ: ใบแรกไปลงโฟลเดอร์สรธรและอนุมัติแล้ว ใบหลังไปลงธีรพงศ์และถูกพักเป็นซ้ำ
j_first, j_second = job("สรธร", "2026-09-07"), job("ธีรพงศ์", "2026-09-07")
first = trip(j_first, "2W-Win ป๋อง = 71_86659+86658_฿30.jpg", committed=1, booking_code="A" * 20,
             net_earnings=30.0, customer_image="WK37-สรธร3.jpg", image_hash="h1")
trip(j_second, "2W-Win ป๋อง = 71_86659+86658_฿30.jpg", status="duplicate", booking_code="A" * 20,
     net_earnings=30.0, duplicate_of=first, image_hash="h1",
     note="ทิ้งอัตโนมัติ: ซ้ำกับ ... ที่อนุมัติแล้ว")
# รหัสเดียวกันแต่ยอดไม่ตรง ยังอยู่ในคิว ต้องให้คนตัดสิน ห้ามส่งไปตีกลับ
j_third = job("ชยพล", "2026-09-07")
trip(j_third, "2W-Win เอก = 306_87733+877321_฿53.jpg", booking_code="B" * 20, net_earnings=53.0,
     duplicate_of=first)

_rows, _twins = rep.load("2026-09-07", "2026-09-13")
_cases, _undecided = rep.build(_rows, _twins)
check("เคสที่พร้อมตีกลับมี 1 ใบ", len(_cases) == 1)
check("เคสที่ต้องให้คนตัดสินไม่ปนเข้าไป", len(_undecided) == 1)
check("ตีกลับไปที่ ป๋อง ไม่ใช่ ธีรพงศ์ ที่ระบบเอารูปไปวาง",
      _cases[0]["ไรเดอร์ผู้ส่ง"] == "2W-Win ป๋อง = 71" and _cases[0]["โฟลเดอร์ที่ระบบวาง"] == "ธีรพงศ์")
check("บอกได้ว่าซ้ำกับไฟล์ไหน และรูปที่ส่งลูกค้าชื่ออะไร",
      _cases[0]["ซ้ำกับไฟล์"] == "2W-Win ป๋อง = 71_86659+86658_฿30.jpg"
      and _cases[0]["รูปที่ส่งลูกค้า"] == "WK37-สรธร3.jpg")
check("ไฟล์ไบต์เดียวกัน ระบุประเภทว่าไฟล์เดิมส่งซ้ำ",
      _cases[0]["ประเภทการซ้ำ"] == rep.KIND_SAME_FILE)
check("สรุปรวมรายผู้ส่งได้", rep.summary(_cases)["2W-Win ป๋อง = 71"]["ใบซ้ำ"] == 1)

_x = rep.write_xlsx(os.path.join(WORK, "dup.xlsx"), _cases, _undecided)
from openpyxl import load_workbook                              # noqa: E402
_wb = load_workbook(_x)
check("ไฟล์ Excel มีห้าชีต (สามชีตเดิม + ซ้ำก่อนอ่าน 2 ชีต)", _wb.sheetnames == ["สรุปรายผู้ส่ง", "รายใบ", "ค้างตัดสิน", "ซ้ำก่อนอ่าน-รายอัลบั้ม", "ซ้ำก่อนอ่าน-รายใบ"])
check("ชีตรายใบมีช่องไว้กรอกผลการตีกลับ",
      {"สถานะเคส", "ผู้ตีกลับ", "คำตอบผู้ส่ง"} <= {c.value for c in _wb["รายใบ"][1]})
check("รันทั้งสคริปต์แล้วไม่ error", rep.main(["--from", "2026-09-07", "--to", "2026-09-13"]) == 0)

# วางลง Drive แยกจากไฟล์หลักที่ส่งลูกค้า — Ops 2026-09-13
check("ชื่อไฟล์บอกสัปดาห์ ไม่ใช่วันที่", rep.report_name("2026-09-07") == "รูปซ้ำ_2026-W37.xlsx")


class _FakeDrive:
    def __init__(self):
        self.folders, self.files = {}, {}

    def ensure_folder(self, parent, name):
        return self.folders.setdefault((parent, name), f"{parent}/{name}")

    def upload_xlsx(self, parent, name, data):
        self.files[(parent, name)] = data            # ชื่อเดิมในโฟลเดอร์เดิม = ทับ
        return f"id:{name}"


_d = _FakeDrive()
_data = rep.build_xlsx(_cases, _undecided)
_folder, _fid = rep.upload_to_drive(_d, "EXPORTS", "2026-09-07", _data)
check("วางไว้ในโฟลเดอร์ของตัวเองใต้ Exports ไม่ใช่ปนกับไฟล์หลัก",
      _folder == f"EXPORTS/{rep.DRIVE_DIR}" and ("EXPORTS", "Rider Trips.xlsx") not in _d.files)
check("ไฟล์ที่วางคือ Excel ที่เปิดได้จริง",
      load_workbook(__import__("io").BytesIO(_d.files[(_folder, "รูปซ้ำ_2026-W37.xlsx")])).sheetnames
      == ["สรุปรายผู้ส่ง", "รายใบ", "ค้างตัดสิน", "ซ้ำก่อนอ่าน-รายอัลบั้ม", "ซ้ำก่อนอ่าน-รายใบ"])
rep.upload_to_drive(_d, "EXPORTS", "2026-09-07", _data)
check("รันซ้ำแล้วทับฉบับเดิม ไม่สร้างไฟล์ชื่อซ้ำเพิ่ม", len(_d.files) == 1)

# ท้ายรอบ ingest เรียก refresh_week — ต้องได้ผลเหมือนกดปุ่มเอง
_d2, _said = _FakeDrive(), []
_res = rep.refresh_week(_d2, "EXPORTS", "2026-09-07", "2026-09-13", log=_said.append)
check("ท้ายรอบประกอบรายงานแล้ววางลง Drive", _res and _res["cases"] == 1 and len(_d2.files) == 1)
check("บอกใน log ว่าวางไว้ที่ไหน", any("รูปซ้ำ_2026-W37.xlsx" in s and "drive.google.com" in s for s in _said))
_d2b, _said2 = _FakeDrive(), []
_res2 = rep.refresh_week(_d2b, "EXPORTS", "2026-09-07", "2026-09-13", log=_said2.append)
check("❗ รอบถัดไปที่ใบซ้ำไม่เปลี่ยน: ไม่อัปโหลดซ้ำ", _res2.get("unchanged") and not _d2b.files)
check("บอกใน log ว่าไม่ต้องเขียนใหม่", any("ไม่มีอะไรเปลี่ยน" in s for s in _said2))
_d2c = _FakeDrive()
import db as _db                                                # noqa: E402
_db.state_set(f"{rep.FINGERPRINT_KEY}:2026-09-07", "stale")
check("เนื้อหาเปลี่ยน: อัปโหลดใหม่",
      rep.refresh_week(_d2c, "EXPORTS", "2026-09-07", "2026-09-13", log=lambda *a: None)["file_id"]
      and len(_d2c.files) == 1)
_d3 = _FakeDrive()
check("สัปดาห์ที่ไม่มีใบซ้ำ ไม่สร้างไฟล์เปล่า",
      rep.refresh_week(_d3, "EXPORTS", "2026-08-10", "2026-08-16", log=lambda *a: None) is None
      and not _d3.files)
import config as _cfg                                           # noqa: E402
check("ท้ายรอบเปิดใช้เป็นค่าเริ่มต้น", _cfg.DUP_REPORT_IN_ROUND is True)

# ลูกค้า 2026-09-14: ซ้ำข้ามสัปดาห์ไม่นับเป็นซ้ำ — ไม่มีอะไรให้ตีกลับ
j_w36 = job("มารือ", "2026-08-31")
old = trip(j_w36, "LINE_ALBUM_DL-TAE_260828_289.jpg", committed=1, booking_code="C" * 20,
           net_earnings=159.0, image_hash="h9")
j_w37 = job("ฉลองรัฐ", "2026-09-07")
trip(j_w37, "LINE_ALBUM_2WTAEW6=295_260911_127.jpg", status="duplicate", booking_code="C" * 20,
     net_earnings=159.0, note="ทิ้งอัตโนมัติ: ซ้ำกับ ... ที่อนุมัติแล้ว")          # รหัสเดียวกัน คนละสัปดาห์
trip(j_w37, "same-bytes.jpg", status="duplicate", booking_code="D" * 20, net_earnings=40.0,
     duplicate_of=old, image_hash="h9")                                          # ไฟล์เดิมเป๊ะ คนละสัปดาห์
_rows2, _twins2 = rep.load("2026-09-07", "2026-09-13")
_cases2, _und2 = rep.build(_rows2, _twins2)
check("ซ้ำข้ามสัปดาห์ (รหัสเดียวกัน) ไม่อยู่ในรายการตีกลับ",
      all(c["ไฟล์ที่ซ้ำ"] != "LINE_ALBUM_2WTAEW6=295_260911_127.jpg" for c in _cases2 + _und2))
check("ไฟล์เดิมเป๊ะที่ต้นฉบับอยู่สัปดาห์อื่น ก็ไม่อยู่ในรายการตีกลับ",
      all(c["ไฟล์ที่ซ้ำ"] != "same-bytes.jpg" for c in _cases2 + _und2))
check("ซ้ำในสัปดาห์เดียวกันยังอยู่ครบเหมือนเดิม", len(_cases2) == 1 and len(_und2) == 1)
# รหัสเดียวกันทั้งสองสัปดาห์: ต้นฉบับที่ใช้ต้องเป็นของสัปดาห์เดียวกัน ไม่ใช่ใบแรกของโปรเจกต์
trip(job("ก", "2026-08-31"), "w36-first.jpg", committed=1, booking_code="E" * 20, net_earnings=70.0)
j_a = job("ข", "2026-09-07")
trip(j_a, "w37-first.jpg", committed=1, booking_code="E" * 20, net_earnings=70.0)
trip(job("ค", "2026-09-07"), "w37-repeat.jpg", status="duplicate", booking_code="E" * 20, net_earnings=70.0)
_c3, _ = rep.build(*rep.load("2026-09-07", "2026-09-13"))
_rep = [c for c in _c3 if c["ไฟล์ที่ซ้ำ"] == "w37-repeat.jpg"]
check("มีต้นฉบับทั้ง W36 และ W37 → ชี้ต้นฉบับของ W37", len(_rep) == 1 and _rep[0]["ซ้ำกับไฟล์"] == "w37-first.jpg")

# --- ซ้ำก่อนอ่าน: ไฟล์เหมือนกันเป๊ะที่ถูกตัดทิ้งก่อนเสียค่าอ่าน (เฟียส 2026-09-22) ---------------
# W38: 2W-Home Jabzaja ส่งอัลบั้ม 48 งานมาสองชุด — รอบเก่าจดไว้แค่ใน pool_runs จึงต้องเติมย้อนหลัง
import json as _json                                            # noqa: E402
W38 = ("2026-09-14", "2026-09-20")
_dups = [{"week": "Week 14-20 Sep", "group": "2W", "file": f"2W-Home Jabzaja/L_{i + 48}_0.jpg",
          "same_as": f"2W-Home Jabzaja/L_{i}_0.jpg"} for i in range(1, 4)]
_dups.append({"week": "Week 14-20 Sep", "group": "2W", "file": "2W-Home UploadJab/3503_1.jpg",
              "same_as": "2W-Home Jabzaja/L_9_0.jpg"})
_dups.append({"week": "Week 7-13 Sep", "group": "2W", "file": "2W-Win ก/1.jpg", "same_as": "2W-Win ก/2.jpg"})
for _mode in ("move", "move", "report"):       # สำเนาที่ค้างในกองโผล่ซ้ำทุกรอบ · รอบรายงานเฉย ๆ ไม่นับ
    db.record_pool_run(_mode, "Week 14-20 Sep", {}, "", {"started_at": "2026-09-22 13:00:00",
                                                         "duplicates": _dups})
check("เติมย้อนหลังจากผลของรอบก่อน ๆ: ไฟล์เดียวกันนับครั้งเดียว เฉพาะสัปดาห์นี้",
      rep.backfill_pre_read(*W38, log=lambda *a: None) == 4)
check("เติมย้อนหลังครั้งเดียวต่อสัปดาห์", rep.backfill_pre_read(*W38, log=lambda *a: None) == 0)
_pre = rep.pre_read_cases(W38[0])
_kinds = {c["ไฟล์ที่ซ้ำ"]: c["ประเภทการซ้ำ"] for c in _pre}
check("สำเนาในอัลบั้มเดียวกัน กับ ซ้ำข้ามอัลบั้ม แยกกันได้",
      _kinds["L_49_0.jpg"] == "สำเนาในอัลบั้มเดียวกัน" and _kinds["3503_1.jpg"] == "ซ้ำข้ามอัลบั้ม")
db.record_skipped_copies([{"ref": "ingest:X1", "week_from": W38[0], "album": "2W-Home Jabzaja",
                           "file_name": "2W-Home Jabzaja_0+0_฿24.jpg", "drive_id": "X1",
                           "same_as": "2W-Home Jabzaja_0+0_฿24.jpg", "kind": "ส่งซ้ำของที่อ่านไปแล้ว"}] * 2)
_pre = rep.pre_read_cases(W38[0])
check("รอบ ingest จดรูปที่ข้ามไว้ด้วย และจดซ้ำไม่เพิ่ม", len(_pre) == 5)
_s = rep.summary([], _pre)
check("❗ นับเข้า 'ใบซ้ำ' ของคนส่ง (ข้อ 2 ข.) และแยกคอลัมน์ ซ้ำก่อนอ่าน ให้เห็น",
      _s["2W-Home Jabzaja"]["ใบซ้ำ"] == 4 and _s["2W-Home Jabzaja"][rep.KIND_PRE] == 4)
_by = {r[0]: r for r in rep.pre_read_by_album(_pre)}
check("สรุปรายอัลบั้มเป็นประโยคที่ใช้ทักคนส่งได้",
      _by["2W-Home Jabzaja"][1] == 4 and "สำเนาในอัลบั้มเดียวกัน 3 ใบ" in _by["2W-Home Jabzaja"][2])
check("มีลิงก์รูปเมื่อรู้รหัสไฟล์", any(c["ลิงก์รูป"] and "X1" in c["ลิงก์รูป"] for c in _pre))
_d4 = _FakeDrive()
_r4 = rep.refresh_week(_d4, "EXPORTS", *W38, log=lambda *a: None)
_wb4 = load_workbook(__import__("io").BytesIO(next(iter(_d4.files.values()))))
check("❗ สัปดาห์ที่มีแต่ซ้ำก่อนอ่าน ก็ได้รายงาน", _r4 and _r4["pre"] == 5 and _d4.files)
check("ชีตรายอัลบั้มและรายใบมีข้อมูลครบ",
      _wb4["ซ้ำก่อนอ่าน-รายอัลบั้ม"].max_row == 3 and _wb4["ซ้ำก่อนอ่าน-รายใบ"].max_row == 6)
_hdr = [c.value for c in _wb4["สรุปรายผู้ส่ง"][1]]
check("ชีตสรุปรายผู้ส่งมีคอลัมน์ ซ้ำก่อนอ่าน", rep.KIND_PRE in _hdr)

shutil.rmtree(WORK, ignore_errors=True)
print("\nสรุป:", "ผ่านทั้งหมด ✅" if ok else "มีข้อที่ไม่ผ่าน ✗")
sys.exit(0 if ok else 1)
