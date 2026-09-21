# -*- coding: utf-8 -*-
"""A rider's trips taken out of a week and delivered on their own (Ops 2026-09-21: 20 urgent
2 W Saver fix-ups from one rider, out of W38, dated into 24–30 Aug, with their own pictures).

What must hold: the week's count drops by exactly those trips and nothing else moves; the rows
reach no weekly file — not W38's, and not the closed first workbook their new dates fall in;
the batch folder holds the workbook and only its own pictures; nothing is deleted.
"""
import io
import os
import shutil
import sys
import tempfile
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
WORK = tempfile.mkdtemp(prefix="pocr-aside-")
os.environ["PHOTO_OCR_DATA"] = WORK
os.environ.pop("DATABASE_URL", None)
sys.path.insert(0, "backend")

import openpyxl                                                   # noqa: E402
from sqlalchemy import insert                                     # noqa: E402

import db                                                         # noqa: E402
import excel_writer as xw                                         # noqa: E402
import set_aside as sa                                            # noqa: E402
from drive_client import LocalDrive                               # noqa: E402

db.init_db()
ok = True


def check(name, cond):
    global ok
    print(("✅" if cond else "✗"), name)
    ok = ok and cond


# --- a Drive tree like the real one ---------------------------------------------------------
root = Path(WORK) / "drive"
inbox = root / "inbox"
rider_dir = inbox / "Week 14-20 Sep" / "2 W Saver" / "62-สรธร Home"
other_dir = inbox / "Week 14-20 Sep" / "2 W Saver" / "63-อมฤต Win"
exp_grp = inbox / "Export Pic" / "2026-W38" / "2 W Saver"
for d in (rider_dir, other_dir, exp_grp):
    d.mkdir(parents=True)
drive = LocalDrive(root)

W38 = ("2026-09-14", "2026-09-20")
j_me = db.create_job("สรธร Home", "Trips", *W38, category="2 W Saver", drive_folder_id=str(rider_dir))
j_std = db.create_job("สรธร Home", "Trips", *W38, category="2 W Standard")
j_other = db.create_job("อมฤต Win", "Trips", *W38, category="2 W Saver", drive_folder_id=str(other_dir))
with db.engine.begin() as c:
    for i in range(20):
        day = f"2026-09-{14 + i % 7:02d}"
        fn, pic = f"S__{1000 + i}_0.jpg", f"WK38-สรธร Home{i + 1}.jpg"
        (rider_dir / fn).write_bytes(b"stitched" + bytes([i]))
        (exp_grp / pic).write_bytes(b"delivered" + bytes([i]))
        c.execute(insert(db.trips).values(
            job_id=j_me, file_name=fn, status="done", committed=1, service_type="Saver Bike",
            trip_date=day, trip_time=f"{8 + i // 7:02d}:00", customer_image=pic, base_fare=30 + i,
            net_earnings=30 + i, passenger_paid=33 + i, app_fee=-1, passenger_total=32 + i,
            note=("เดิมมีโน้ต" if i == 0 else None)))
    c.execute(insert(db.trips).values(job_id=j_std, file_name="std.jpg", status="done", committed=1,
                                      service_type="Standard Bike", trip_date="2026-09-15"))
    for i in range(21):
        c.execute(insert(db.trips).values(job_id=j_other, file_name=f"o{i}.jpg", status="done",
                                          committed=1, service_type="Saver Bike", trip_date="2026-09-16"))

before = db.week_group_counts(*W38)
picked = sa.picked_rows("สรธร Home", "2 W Saver", *W38)
check("เลือกได้ 20 งานของไรเดอร์คนเดียว กลุ่ม 2 W Saver เท่านั้น", len(picked) == 20
      and all(j["id"] == j_me for _r, j in picked))
check("ไม่หยิบงาน 2 W Standard ของคนเดียวกัน และไม่หยิบของไรเดอร์อื่น",
      all(r["file_name"].startswith("S__") for r, _j in picked))

done, jid = sa.apply(drive, str(inbox), picked, "สรธร Home", "2 W Saver", *W38,
                     "2026-08-24", "2026-08-30", "งานแก้ 2W Saver 24-30 Aug", log=lambda *_: None)

# --- the week ------------------------------------------------------------------------------
after = db.week_group_counts(*W38)
check("W38 2 W Saver ลดลง 20 พอดี", before["2 W Saver"] - after.get("2 W Saver", 0) == 20)
check("กลุ่มอื่นไม่ขยับ", before["2 W Standard"] == after["2 W Standard"])
card = db.week_scorecard(*W38)
check("ใบตรวจ W38 ไม่นับชุดนี้แล้ว", card["by_service"].get("Saver Bike") == 21)
rows = db.query_trips(committed_only=True)
check("❗ ไม่อยู่ในไฟล์ไหนเลย: Phase 3 ของ W38", not any(r["file_name"].startswith("S__") for r in rows))
first = openpyxl.load_workbook(io.BytesIO(xw.build_workbook(rows)))[xw.SHEET]
check("❗ ไม่ไหลเข้าไฟล์แรก (W33–W35) ทั้งที่วันที่ใหม่อยู่ในเดือนสิงหา", first.max_row == 1)

# --- the rows --------------------------------------------------------------------------------
batch = sa.batch_rows(jid)
check("แถวยังอยู่ครบ 20 ไม่มีอะไรถูกลบ", len(batch) == 20)
check("สถานะ set_aside", all(r["status"] == "set_aside" for r in batch))
check("วันที่อยู่ใน 24–30 Aug ทุกแถว", all("2026-08-24" <= r["trip_date"] <= "2026-08-30" for r in batch))
check("กระจายครบ 7 วัน", len({r["trip_date"] for r in batch}) == 7)
check("โน้ตเก็บวันที่บนสลิปและที่มาไว้ และโน้ตเดิมยังอยู่",
      all("วันที่บนสลิป 2026-09-" in r["note"] and f"job #{j_me}" in r["note"] for r in batch)
      and any(r["note"].endswith("เดิมมีโน้ต") for r in batch))
check("ชื่อรูปใหม่เป็นของสัปดาห์ใหม่ WK35-สรธร Home1..20",
      sorted(r["customer_image"] for r in batch) == sorted(f"WK35-สรธร Home{i}.jpg" for i in range(1, 21)))
check("เงินไม่ขยับ", sorted(r["base_fare"] for r in batch) == [30 + i for i in range(20)])

# --- the folders -------------------------------------------------------------------------------
bdir = inbox / "Export Pic" / "งานแก้ 2W Saver 24-30 Aug"
pics = sorted(p.name for p in bdir.glob("*.jpg"))
check("โฟลเดอร์ชุดงานมีรูปส่งลูกค้า 20 รูปพอดี ชื่อใหม่", pics == sorted(f"WK35-สรธร Home{i}.jpg" for i in range(1, 21)))
check("รูปคือภาพเดิมทุกไบต์", all((bdir / f"WK35-สรธร Home{i}.jpg").read_bytes().startswith(b"delivered")
                             for i in range(1, 21)))
check("รูปรวมร่างต้นฉบับย้ายไป _ต้นฉบับ", len(list((bdir / "_ต้นฉบับ").glob("*.jpg"))) == 20
      and not list(rider_dir.glob("*.jpg")))
check("รูปส่งลูกค้าเดิมออกจาก Export Pic ของ W38 ไป _แทนที่แล้ว (ไม่ลบ)",
      not list(exp_grp.glob("*.jpg"))
      and len(list((inbox / "Export Pic" / "2026-W38" / "_แทนที่แล้ว").glob("*.jpg"))) == 20)
check("โฟลเดอร์ไรเดอร์อื่นไม่ถูกแตะ", other_dir.exists())
xl = bdir / "Rider Trips งานแก้ 2W Saver 24-30 Aug.xlsx"
check("ไฟล์ Excel อยู่ในโฟลเดอร์เดียวกับรูป", xl.exists())
ws = openpyxl.load_workbook(xl)[xw.LOCATION_SHEET]
H = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
check("ไฟล์มี 20 แถว คอลัมน์แบบ Phase 3", ws.max_row == 21
      and [ws.cell(1, c).value for c in range(1, ws.max_column + 1)] == xw.FARE_LINES_HEADERS)
check("ช่อง Image ในไฟล์ตรงกับชื่อรูปในโฟลเดอร์",
      sorted(ws.cell(r, H["Image"]).value for r in range(2, 22)) == pics)
check("วันที่ในไฟล์เป็นเดือนสิงหา", all(ws.cell(r, H["Date & Time"]).value.strftime("%Y-%m") == "2026-08"
                                   for r in range(2, 22)))

# --- run again ---------------------------------------------------------------------------------
check("รันซ้ำ: ไม่มีงานให้ดึงแล้ว", sa.picked_rows("สรธร Home", "2 W Saver", *W38) == [])

shutil.rmtree(WORK, ignore_errors=True)
print("\nสรุป:", "ผ่านทั้งหมด ✅" if ok else "มีข้อที่ไม่ผ่าน ✗")
sys.exit(0 if ok else 1)
