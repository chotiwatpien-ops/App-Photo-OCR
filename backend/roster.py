# -*- coding: utf-8 -*-
"""Build next week's rider folders on Drive from Ops' name list.

Ops keeps a pool of names per wheel count and wants them rotated: most of a week's riders carry
over from the week before, the rest are swapped for names that sat out. The label that says what
kind of driver it is — Win, Home, Taxi — travels with the name, into the folder and from there
into the workbook's Driver Name column, because 'ตฤณ Win' and 'ตฤณ Taxi' are two different people.

    python backend/roster.py --import "Name list by Type (Rider Project).xlsx"
    python backend/roster.py --week "Week 7-13 Sep" --dry-run
    python backend/roster.py --week "Week 7-13 Sep"

Rotation is a request, not a promise: a group can only take in as many new names as the pool has
spare. With 145 names for 140 places the 2W side can turn over 5 a week however loudly 20% is
asked for, so the plan reports what it actually managed.
"""
import argparse
import random
import re
import sys

import config
import db
from drive_client import DriveClient

# Ops' sheets, and the label each one puts after the name
SHEETS = {
    "2 W (Win)": ("2W", "Win"),
    "2 W Not Win": ("2W", "Home"),
    "4 W Taxi": ("4W", "Taxi"),
    "4 W Not Taxi": ("4W", "Home"),
}
# which pool a vehicle group draws from
GROUP_WHEEL = {"2 W Saver": "2W", "2 W Standard": "2W",
               "4 W Standard": "4W", "4 W Saver": "4W"}
PER_GROUP = 70                      # 1,470 trips a week ÷ 21 a rider
KEEP = 0.8                          # how much of last week to carry over, when the pool allows


def log(msg):
    print(msg, flush=True)


def label(name, kind):
    """'ตฤณ' + 'Taxi' -> 'ตฤณ Taxi'. Home riders carry their label too, so that a name in both
    the Not-Win and the Not-Taxi list does not become one person in the workbook."""
    return f"{name} {kind}".strip()


def _drive():
    """Actions passes the token in the environment; a laptop keeps it in a file next to the code."""
    import os
    tok = None if os.environ.get("DRIVE_OAUTH_TOKEN_JSON") else "drive_token.json"
    return DriveClient(oauth_token_json=tok)


def _list_all(drive, parent_id):
    """Everything in a folder, pictures or not — list_images filters to images."""
    return drive._list(f"'{parent_id}' in parents and trashed=false", "id, name")


def import_pool(path) -> int:
    """path: a local .xlsx, or bytes already fetched from Drive."""
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(path) if isinstance(path, (bytes, bytearray)) else path,
                                read_only=True)
    rows = []
    for sheet, (wheel, kind) in SHEETS.items():
        if sheet not in wb.sheetnames:
            log(f"⚠ ไม่มีชีต {sheet!r} ในไฟล์ — ข้าม")
            continue
        for (v,) in wb[sheet].iter_rows(min_row=2, max_col=1, values_only=True):
            if v and str(v).strip():
                rows.append((str(v).strip(), wheel, kind))
    n = db.name_pool_load(rows)
    for wheel in ("2W", "4W"):
        log(f"  {wheel}: {sum(1 for _, w, _ in rows if w == wheel)} ชื่อ")
    return n


def read_week(drive, inbox_id, week_name):
    """{group: [rider folder names]} for a week already on Drive, or {} if it is not there."""
    wk = next((f for f in drive.list_folders(inbox_id) if f["name"].strip() == week_name.strip()), None)
    if not wk:
        return {}
    out = {}
    for cat in drive.list_folders(wk["id"]):
        if cat["name"].strip().lower() in ("pool", "กอง"):
            continue
        kids = drive.list_folders(cat["id"])
        names = []
        for k in kids:
            if re.match(r"^\s*admin\b", k["name"], re.IGNORECASE):
                names += [x["name"] for x in drive.list_folders(k["id"])]
            else:
                names.append(k["name"])
        out[cat["name"]] = names
    return out


NUM_PREFIX = re.compile(r"^\s*\d{1,3}\s*[-. ]?\s*")


def bare(folder_name):
    """'01-ตฤณ Taxi ' -> 'ตฤณ Taxi'."""
    return re.sub(r"\s+", " ", NUM_PREFIX.sub("", folder_name or "")).strip()


def plan(previous, pools, per_group=PER_GROUP, keep=KEEP, seed=None):
    """What each group's rider list should be next week.

    previous: {group: [names as they appear now]}   pools: {wheel: [(name, kind)]}
    Returns ({group: [labelled names]}, [notes]). A name is never used twice in one week, not
    even across groups: a job is keyed on rider and week, so two folders with one name are one
    rider to everything downstream."""
    rng = random.Random(seed)
    notes, chosen = {}, {}
    labelled = {w: [label(n, k) for n, k in p] for w, p in pools.items()}
    # Everyone who worked last week, whatever group they were in. A name dropped from one group
    # must not reappear in another the same week and be counted as a new face — that is a
    # reshuffle, not a rotation, and it is what makes 'turn over 20%' impossible when the pool
    # is barely larger than the roster: only the names that sat out are genuinely new.
    last_week = {bare(x) for names in (previous or {}).values() for x in names}
    groups = [g for g in sorted(previous or GROUP_WHEEL, key=str)]
    for g in groups:
        if g not in GROUP_WHEEL:
            notes[g] = "ไม่รู้ว่ากลุ่มนี้ดึงชื่อจากพูลไหน — ข้าม"

    # Share out the new faces before filling any group. Taken group by group instead, the first
    # one served empties the pool and the last never rotates at all: with 145 names for 140
    # places, 2 W Saver took all five and 2 W Standard stood still week after week.
    want_new = int(round(per_group * (1 - keep)))
    share, free = {}, {}
    for wheel in {GROUP_WHEEL[g] for g in groups if g in GROUP_WHEEL}:
        mine = [g for g in groups if GROUP_WHEEL.get(g) == wheel]
        free[wheel] = [n for n in labelled.get(wheel, []) if n not in last_week]
        rng.shuffle(free[wheel])
        budget = min(len(free[wheel]), want_new * len(mine))
        for i, g in enumerate(mine):                       # spread the remainder over the first few
            share[g] = budget // len(mine) + (1 if i < budget % len(mine) else 0)

    taken = set()
    for group in groups:
        wheel = GROUP_WHEEL.get(group)
        if not wheel:
            continue
        old = [bare(x) for x in (previous or {}).get(group, [])]
        old = [n for n in dict.fromkeys(old) if n in labelled.get(wheel, []) and n not in taken]
        new = [n for n in free[wheel] if n not in taken][:share.get(group, 0)]
        kept = rng.sample(old, min(len(old), per_group - len(new))) if old else []
        if len(kept) + len(new) < per_group:
            # short even after keeping everyone: borrow whoever is still unspoken for
            spare = [n for n in labelled.get(wheel, [])
                     if n not in taken and n not in kept and n not in new]
            rng.shuffle(spare)
            new += spare[:per_group - len(kept) - len(new)]
        taken.update(kept + new)
        chosen[group] = sorted(kept + new)
        got = len(chosen[group])
        fresh = len([n for n in chosen[group] if n not in last_week])
        notes[group] = (f"{got} คน · เดิม {got - fresh} · ใหม่ {fresh} "
                        f"({fresh / got * 100:.0f}%)" if got else "0 คน")
        if got < per_group:
            notes[group] += f" · ⚠ ขาด {per_group - got} (พูลหมด)"
    return chosen, notes


def folder_name(i, name):
    return f"{i:02d}-{name}"


def apply(drive, inbox_id, week_name, chosen, dry_run=True):
    made = 0
    wk = next((f for f in drive.list_folders(inbox_id) if f["name"].strip() == week_name.strip()), None)
    if not wk and dry_run:
        log(f"(ยังไม่มีโฟลเดอร์สัปดาห์ {week_name!r} — จะสร้างใหม่)")
    elif not wk:
        wk = {"id": drive.ensure_folder(inbox_id, week_name), "name": week_name}
    for group, names in sorted(chosen.items()):
        # what is already there stays: this can be run again after Ops adds a name by hand
        cat = existing = None
        if wk:
            cat = next((f for f in drive.list_folders(wk["id"]) if f["name"].strip() == group), None)
            if cat is None and not dry_run:
                cat = {"id": drive.ensure_folder(wk["id"], group)}
            existing = {bare(f["name"]) for f in drive.list_folders(cat["id"])} if cat else set()
        have = existing or set()
        todo = [n for n in names if n not in have]
        if not dry_run:
            for i, n in enumerate(names, 1):
                if n not in have:
                    drive.ensure_folder(cat["id"], folder_name(i, n))
        made += len(todo)
        log(f"  {group}: {'จะสร้าง' if dry_run else 'สร้างแล้ว'} {len(todo)} โฟลเดอร์"
            + (f" · มีอยู่แล้ว {len(have & set(names))}" if have else ""))
    return made


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--import", dest="imp", help="โหลดรายชื่อจากไฟล์ Excel ในเครื่อง")
    ap.add_argument("--import-drive", dest="imp_drive", default="",
                    help="โหลดรายชื่อจากไฟล์ Excel ที่วางไว้ใน Inbox บน Drive (ใส่ชื่อไฟล์)")
    ap.add_argument("--week", help="ชื่อโฟลเดอร์สัปดาห์ที่จะสร้าง เช่น 'Week 7-13 Sep'")
    ap.add_argument("--from", dest="prev", default="", help="สัปดาห์ที่เอาชื่อเดิมมา (ว่าง = สัปดาห์ล่าสุดที่มี)")
    ap.add_argument("--per-group", type=int, default=PER_GROUP)
    ap.add_argument("--keep", type=float, default=KEEP, help="สัดส่วนชื่อเดิมที่เก็บไว้ (0.8 = 80%%)")
    ap.add_argument("--seed", type=int, default=None, help="ล็อกการสุ่มให้ได้ผลเดิม (ทดสอบ)")
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    db.init_db()

    if a.imp or a.imp_drive:
        if a.imp_drive:
            drive = _drive()
            f = next((x for x in drive.list_images(config.DRIVE_INBOX_FOLDER_ID)
                      if x["name"].strip() == a.imp_drive.strip()), None)
            if f is None:                       # list_images only yields pictures; look wider
                f = next((x for x in _list_all(drive, config.DRIVE_INBOX_FOLDER_ID)
                          if x["name"].strip() == a.imp_drive.strip()), None)
            if f is None:
                log(f"✗ ไม่พบไฟล์ {a.imp_drive!r} ใน Inbox บน Drive")
                return 1
            log(f"นำเข้ารายชื่อจาก Drive: {f['name']}")
            src = drive.download(f["id"])
        else:
            log(f"นำเข้ารายชื่อจาก {a.imp}")
            src = a.imp
        log(f"เก็บลงฐานข้อมูลแล้ว {import_pool(src)} ชื่อ")
        if not a.week:
            return 0

    if not a.week:
        ap.error("ต้องระบุ --week หรือ --import")

    pools = {w: db.name_pool_for(w) for w in ("2W", "4W")}
    if not any(pools.values()):
        log("✗ ยังไม่มีรายชื่อในฐานข้อมูล — สั่ง --import ก่อน")
        return 1
    drive = _drive()
    inbox = config.DRIVE_INBOX_FOLDER_ID
    prev_name = a.prev
    if not prev_name:
        weeks = sorted((f["name"] for f in drive.list_folders(inbox)
                        if f["name"].lower().startswith("week") and not config.week_ignored(f["name"])),
                       key=lambda s: s)
        prev_name = next((w for w in reversed(weeks) if w.strip() != a.week.strip()), "")
    log(f"เอาชื่อเดิมจาก: {prev_name or '(ไม่มี — สุ่มใหม่ทั้งหมด)'}")
    previous = read_week(drive, inbox, prev_name) if prev_name else {}
    for g, v in sorted(previous.items()):
        log(f"  {g}: {len(v)} คน")

    chosen, notes = plan(previous or {g: [] for g in GROUP_WHEEL if g != "4 W Saver"},
                         pools, a.per_group, a.keep, a.seed)
    log(f"\nแผนสำหรับ {a.week}:")
    for g in sorted(chosen):
        log(f"  {g}: {notes[g]}")
    n = apply(drive, inbox, a.week, chosen, dry_run=a.dry_run)
    log(f"\n{'(ทดลอง) จะสร้าง' if a.dry_run else 'สร้าง'} {n} โฟลเดอร์")
    return 0


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.exit(main())
